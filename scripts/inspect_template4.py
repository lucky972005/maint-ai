import docx

doc = docx.Document(r"C:\Users\hetjo\Downloads\Internship_UDP_Training_Report_Template for reporting 4.docx")

with open(r"C:\dev\pro\template4_full_inspection.txt", "w", encoding="utf-8") as out:
    out.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            out.write(f"P[{i}]: {text}\n")

    out.write("\n=== TABLES ===\n")
    for i, t in enumerate(doc.tables):
        out.write(f"\n--- TABLE {i} ({len(t.rows)}x{len(t.columns)}) ---\n")
        for r_i, r in enumerate(t.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            out.write(f"  Row {r_i}: {cells}\n")

print("Dumped template4_full_inspection.txt")
