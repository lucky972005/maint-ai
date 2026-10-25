import docx
import zipfile
import os

docx_path = r"C:\dev\pro\MaintAI_Reporting_4_Het_Joshi.docx"
print("File size:", os.path.getsize(docx_path), "bytes")

doc = docx.Document(docx_path)
print("Paragraphs count:", len(doc.paragraphs))
print("Tables count:", len(doc.tables))

print("\n--- CHECKING TABLES ---")
for i, t in enumerate(doc.tables):
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols")
    for r_i, r in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ')[:40] for c in r.cells]
        print(f"  r{r_i}: {cells}")

print("\n--- CHECKING EMBEDDED MEDIA IN DOCX ZIP ---")
with zipfile.ZipFile(docx_path, "r") as z:
    for name in z.namelist():
        if name.startswith("word/media/"):
            info = z.getinfo(name)
            print(f"  {name}: {info.file_size} bytes")

print("\n--- CHECKING KEY PARAGRAPHS ---")
check_indices = [0, 1, 6, 7, 9, 10, 11, 12, 13, 14, 15, 17, 18, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40, 41, 43, 44, 45, 47, 48, 49, 51, 52, 54, 55, 57, 59, 60, 62, 63, 65, 66, 68, 69, 79, 80, 81, 82, 83, 84, 93, 96, 98, 99]

with open(r"C:\dev\pro\report4_verification.txt", "w", encoding="utf-8") as f:
    for i in check_indices:
        if i < len(doc.paragraphs):
            p = doc.paragraphs[i]
            f.write(f"P[{i}]: {p.text[:120]}\n")
            for r in p.runs:
                if "drawing" in r._r.xml:
                    f.write(f"    [RUN HAS IMAGE/DRAWING: {r.text}]\n")

print("Dumped report4_verification.txt")
