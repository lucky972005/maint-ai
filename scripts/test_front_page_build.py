import docx
from docx.shared import Pt
import win32com.client
import fitz
import os

TEMPLATE_PATH = r"C:\Users\hetjo\Downloads\report\Project Report Front page for Hard Binding.docx"

def test_with_title():
    doc = docx.Document(TEMPLATE_PATH)
    
    # In doc.paragraphs:
    # P00: TRAINING / INTERNSHIP / USER-DEFINED PROJECT
    # P01: A PROJECT REPORT
    # P02: ''
    # Let's replace P02 with Project Title
    p2 = doc.paragraphs[2]
    p2.text = "“MaintAI — An AI-Powered Progressive Web Application for Industrial Maintenance Assistance”"
    p2.runs[0].font.size = Pt(13)
    p2.runs[0].bold = True
    p2.paragraph_format.space_before = Pt(4)
    p2.paragraph_format.space_after = Pt(8)

    for p in doc.paragraphs:
        if '[BRANCH / SPECIALISATION]' in p.text:
            p.text = p.text.replace('[BRANCH / SPECIALISATION]', 'COMPUTER ENGINEERING')
            p.runs[0].font.size = Pt(14)
            p.runs[0].bold = True
        if '[STUDENT NAME]' in p.text:
            p.text = p.text.replace('[STUDENT NAME]', 'HET JOSHI')
            p.runs[0].font.size = Pt(15)
            p.runs[0].bold = True
        if '[ENROLLMENT NUMBER]' in p.text:
            p.text = p.text.replace('[ENROLLMENT NUMBER]', '24SE09CE003')
            p.runs[0].font.size = Pt(11)
        if '[ORGANISATION NAME]' in p.text:
            p.text = p.text.replace('[ORGANISATION NAME]', 'P P SAVANI UNIVERSITY')
            p.runs[0].font.size = Pt(11)
            p.runs[0].bold = True
        if '[FACULTY MENTOR NAME]' in p.text:
            p.text = p.text.replace('[FACULTY MENTOR NAME]', 'Mr. Ayush Dodia')
            p.runs[0].font.size = Pt(13)
            p.runs[0].bold = True
        if '[Designation, School of Engineering]' in p.text:
            p.text = p.text.replace('[Designation, School of Engineering]', 'Assistant Professor, School of Engineering')
            p.runs[0].font.size = Pt(11)
            p.runs[0].italic = True
        if '[MONTH YEAR]' in p.text:
            p.text = p.text.replace('[MONTH YEAR]', 'OCTOBER 2026')
            p.runs[0].font.size = Pt(12)
            p.runs[0].bold = True

    # Adjust some empty paragraph spacings if needed to ensure 1 page
    # P05, P09, P13, P15, P19, P21 are empty
    for i in [5, 9, 13, 15, 19, 21]:
        doc.paragraphs[i].paragraph_format.space_after = Pt(4)

    out_docx = r"C:\dev\pro\test_front_with_title.docx"
    out_pdf = r"C:\dev\pro\test_front_with_title.pdf"
    doc.save(out_docx)
    
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    doc_w = word.Documents.Open(out_docx, ReadOnly=True)
    doc_w.SaveAs(out_pdf, FileFormat=17)
    doc_w.Close(False)
    word.Quit()
    
    doc_p = fitz.open(out_pdf)
    print("With title Page count:", len(doc_p))

test_with_title()
