import docx

doc = docx.Document(r"C:\Users\hetjo\Downloads\Internship_UDP_Training_Report_Template for reporting 4.docx")

with open(r"C:\dev\pro\body_order.txt", "w", encoding="utf-8") as out:
    p_idx = 0
    t_idx = 0
    for child in doc._element.body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = doc.paragraphs[p_idx]
            txt = p.text.strip().replace('\n', ' ')
            out.write(f"P[{p_idx}]: {txt[:70]}\n")
            p_idx += 1
        elif tag == 'tbl':
            t = doc.tables[t_idx]
            out.write(f"--> TABLE [{t_idx}] ({len(t.rows)}x{len(t.columns)})\n")
            t_idx += 1
        elif tag == 'sectPr':
            out.write("--> Section Properties\n")

print("Dumped body_order.txt")
