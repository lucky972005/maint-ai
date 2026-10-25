import docx

doc = docx.Document(r"C:\dev\pro\MaintAI_Final_Hard_Bound_Project_Report_Het_Joshi.docx")
print("Total tables in docx:", len(doc.tables))
for i in range(min(7, len(doc.tables))):
    t = doc.tables[i]
    first_row = [c.text.strip().replace('\n', ' ')[:25] for c in t.rows[0].cells]
    print(f"Table {i}: {len(t.rows)} rows, {len(t.columns)} cols | Row 0: {first_row}")

print("\n--- Paragraphs mentioning Abbreviations ---")
for p_idx, p in enumerate(doc.paragraphs[:70]):
    if 'abbreviation' in p.text.lower():
        print(f"Para {p_idx}: '{p.text}'")

# Also check template document
tmpl = docx.Document(r"C:\Users\hetjo\Downloads\PPSU_Final_Hard_Bound_Project_Report_Format.docx")
print("\n--- Template Paragraphs mentioning Abbreviations ---")
for p_idx, p in enumerate(tmpl.paragraphs[:70]):
    if 'abbreviation' in p.text.lower():
        print(f"Tmpl Para {p_idx}: '{p.text}'")

print("\n--- Template Tables ---")
for i in range(min(7, len(tmpl.tables))):
    t = tmpl.tables[i]
    first_row = [c.text.strip().replace('\n', ' ')[:25] for c in t.rows[0].cells]
    print(f"Tmpl Table {i}: {len(t.rows)} rows, {len(t.columns)} cols | Row 0: {first_row}")
