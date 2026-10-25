import os
import copy
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
import win32com.client
import fitz

SOURCE_DOCX = r"c:\dev\pro\MaintAI_Final_Hard_Bound_Project_Report_Het_Joshi.docx"
TEST_OUTPUT_DOCX = r"c:\dev\pro\test_combined_report.docx"
TEST_OUTPUT_PDF = r"c:\dev\pro\test_combined_report.pdf"
LOGO_PATH = r"c:\dev\pro\ppsu_naac_front_logo.png"

def build_test():
    doc = docx.Document(SOURCE_DOCX)
    body = doc._body._element

    # Create front page paragraphs
    # We will insert them before the very first element in body!
    first_elem = body[0]

    def make_para(text, font_size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            r = p.add_run(text)
            r.font.name = "Cambria"
            r.font.size = Pt(font_size)
            r.bold = bold
            r.italic = italic
        return p

    front_paras = []

    # P00
    p0 = make_para("TRAINING / INTERNSHIP / USER-DEFINED PROJECT", font_size=18, bold=True, space_before=16, space_after=10)
    front_paras.append(p0)

    # P01
    p1 = make_para("A PROJECT REPORT", font_size=15, bold=True, space_before=0, space_after=6)
    front_paras.append(p1)

    # P02: Project Title
    p2 = make_para("“MaintAI — An AI-Powered Progressive Web Application for Industrial Maintenance Assistance”", font_size=13, bold=True, space_before=4, space_after=8)
    front_paras.append(p2)

    # P03
    p3 = make_para("Submitted in partial fulfilment of the requirements", font_size=11, italic=True, space_before=0, space_after=1.5)
    front_paras.append(p3)

    # P04
    p4 = make_para("for the award of the degree of", font_size=11, italic=True, space_before=0, space_after=8)
    front_paras.append(p4)

    # P05: Spacing
    p5 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p5)

    # P06
    p6 = make_para("BACHELOR OF TECHNOLOGY", font_size=16, bold=True, space_before=0, space_after=1.5)
    front_paras.append(p6)

    # P07
    p7 = make_para("in", font_size=12, bold=True, space_before=0, space_after=1.5)
    front_paras.append(p7)

    # P08
    p8 = make_para("COMPUTER ENGINEERING", font_size=14, bold=True, space_before=0, space_after=8)
    front_paras.append(p8)

    # P09: Spacing
    p9 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p9)

    # P10
    p10 = make_para("Submitted by", font_size=11, italic=True, space_before=0, space_after=2)
    front_paras.append(p10)

    # P11
    p11 = make_para("HET JOSHI", font_size=15, bold=True, space_before=0, space_after=1.5)
    front_paras.append(p11)

    # P12
    p12 = make_para("Enrollment No.: 24SE09CE003", font_size=11, bold=False, space_before=0, space_after=8)
    front_paras.append(p12)

    # P13: Spacing
    p13 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p13)

    # P14
    p14 = make_para("Training / Internship / UDP Organisation: P P SAVANI UNIVERSITY", font_size=11, bold=True, space_before=0, space_after=8)
    front_paras.append(p14)

    # P15: Spacing
    p15 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p15)

    # P16
    p16 = make_para("Under the guidance of", font_size=11, italic=True, space_before=0, space_after=2)
    front_paras.append(p16)

    # P17
    p17 = make_para("Mr. Ayush Dodia", font_size=13, bold=True, space_before=0, space_after=1.5)
    front_paras.append(p17)

    # P18
    p18 = make_para("Assistant Professor, School of Engineering", font_size=11, italic=True, space_before=0, space_after=6)
    front_paras.append(p18)

    # P19: Spacing
    p19 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p19)

    # P20: Logo Image
    p20 = make_para("", font_size=1, space_before=0, space_after=4)
    run_img = p20.add_run()
    run_img.add_picture(LOGO_PATH, width=Inches(2.55))
    front_paras.append(p20)

    # P21: Spacing
    p21 = make_para("", font_size=6, space_before=0, space_after=2)
    front_paras.append(p21)

    # P22
    p22 = make_para("P P SAVANI UNIVERSITY", font_size=16, bold=True, space_before=0, space_after=1)
    front_paras.append(p22)

    # P23
    p23 = make_para("SCHOOL OF ENGINEERING", font_size=12, bold=True, space_before=0, space_after=1.5)
    front_paras.append(p23)

    # P24
    p24 = make_para("P P SAVANI UNIVERSITY, SURAT, GUJARAT", font_size=11, bold=True, space_before=0, space_after=4)
    front_paras.append(p24)

    # P25
    p25 = make_para("OCTOBER 2026", font_size=12, bold=True, space_before=0, space_after=0)
    front_paras.append(p25)

    # Now add sectPr to p25 (End of Front Cover Section)
    sectPr_front = parse_xml(f'''
    <w:sectPr {nsdecls("w", "r")}>
      <w:pgSz w:w="12240" w:h="15840"/>
      <w:pgMar w:top="936" w:right="1152" w:bottom="936" w:left="1152" w:header="720" w:footer="720" w:gutter="0"/>
      <w:pgBorders w:offsetFrom="page" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:top w:val="single" w:sz="12" w:space="18" w:color="B08D35"/>
        <w:left w:val="single" w:sz="12" w:space="18" w:color="B08D35"/>
        <w:bottom w:val="single" w:sz="12" w:space="18" w:color="B08D35"/>
        <w:right w:val="single" w:sz="12" w:space="18" w:color="B08D35"/>
      </w:pgBorders>
      <w:cols w:space="720"/>
      <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')
    p25._p.get_or_add_pPr().append(sectPr_front)

    # Now move all front_paras to the very beginning of body (before first_elem)
    for p in front_paras:
        elem = p._p
        body.remove(elem)
        first_elem.addprevious(elem)

    # Save and test
    doc.save(TEST_OUTPUT_DOCX)
    print(f"Saved {TEST_OUTPUT_DOCX}")

    # Export to PDF via Word
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    try:
        doc_w = word.Documents.Open(os.path.abspath(TEST_OUTPUT_DOCX), ReadOnly=True)
        doc_w.SaveAs(os.path.abspath(TEST_OUTPUT_PDF), FileFormat=17)
        doc_w.Close(False)
        print(f"Exported {TEST_OUTPUT_PDF}")
    finally:
        word.Quit()

    # Inspect PDF pages
    doc_p = fitz.open(TEST_OUTPUT_PDF)
    print(f"Total Pages in Combined PDF: {len(doc_p)}")
    print("=== Page 1 (Front Cover) ===")
    print(doc_p[0].get_text()[:300])
    print("=== Page 2 (Particulars) ===")
    print(doc_p[1].get_text()[:300])
    print("=== Page 3 (Certificate) ===")
    print(doc_p[2].get_text()[:300])

if __name__ == "__main__":
    build_test()
