import docx

doc = docx.Document(r"C:\dev\pro\MaintAI_Final_Hard_Bound_Project_Report_Het_Joshi.docx")
print("Doc paragraphs 45 to 60:")
for i in range(45, min(60, len(doc.paragraphs))):
    print(f"  Para {i}: '{doc.paragraphs[i].text}'")

tmpl = docx.Document(r"C:\Users\hetjo\Downloads\PPSU_Final_Hard_Bound_Project_Report_Format.docx")
print("\nTemplate paragraphs 45 to 60:")
for i in range(45, min(60, len(tmpl.paragraphs))):
    print(f"  Tmpl Para {i}: '{tmpl.paragraphs[i].text}'")
