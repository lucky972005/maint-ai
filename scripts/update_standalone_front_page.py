import os
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client
import fitz

TEMPLATE = r"C:\Users\hetjo\Downloads\report\Project Report Front page for Hard Binding.docx"
OUT_DOCX = r"C:\Users\hetjo\Downloads\report\Project Report Front page for Hard Binding_Het_Joshi.docx"
OUT_PDF = r"C:\Users\hetjo\Downloads\report\Project Report Front page for Hard Binding_Het_Joshi.pdf"

doc = docx.Document(TEMPLATE)
p2 = doc.paragraphs[2]
p2.text = "“MaintAI — An AI-Powered Progressive Web Application for Industrial Maintenance Assistance”"
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
if len(p2.runs) > 0:
    p2.runs[0].font.name = "Cambria"
    p2.runs[0].font.size = Pt(13)
    p2.runs[0].bold = True
p2.paragraph_format.space_before = Pt(4)
p2.paragraph_format.space_after = Pt(8)

for p in doc.paragraphs:
    if '[BRANCH / SPECIALISATION]' in p.text:
        p.text = p.text.replace('[BRANCH / SPECIALISATION]', 'COMPUTER ENGINEERING')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(14)
            p.runs[0].bold = True
    if '[STUDENT NAME]' in p.text:
        p.text = p.text.replace('[STUDENT NAME]', 'HET JOSHI')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(15)
            p.runs[0].bold = True
    if '[ENROLLMENT NUMBER]' in p.text:
        p.text = p.text.replace('[ENROLLMENT NUMBER]', '24SE09CE003')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(11)
    if '[ORGANISATION NAME]' in p.text:
        p.text = p.text.replace('[ORGANISATION NAME]', 'P P SAVANI UNIVERSITY')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(11)
            p.runs[0].bold = True
    if '[FACULTY MENTOR NAME]' in p.text:
        p.text = p.text.replace('[FACULTY MENTOR NAME]', 'Mr. Ayush Dodia')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(13)
            p.runs[0].bold = True
    if '[Designation, School of Engineering]' in p.text:
        p.text = p.text.replace('[Designation, School of Engineering]', 'Assistant Professor, School of Engineering')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(11)
            p.runs[0].italic = True
    if '[MONTH YEAR]' in p.text:
        p.text = p.text.replace('[MONTH YEAR]', 'OCTOBER 2026')
        if len(p.runs) > 0:
            p.runs[0].font.name = "Cambria"
            p.runs[0].font.size = Pt(12)
            p.runs[0].bold = True

for i in [5, 9, 13, 15, 19, 21]:
    doc.paragraphs[i].paragraph_format.space_after = Pt(4)

doc.save(OUT_DOCX)
print("Saved standalone docx:", OUT_DOCX)

word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False
try:
    doc_w = word.Documents.Open(os.path.abspath(OUT_DOCX), ReadOnly=True)
    doc_w.SaveAs(os.path.abspath(OUT_PDF), FileFormat=17)
    doc_w.Close(False)
    print("Exported standalone pdf:", OUT_PDF)
finally:
    word.Quit()

doc_p = fitz.open(OUT_PDF)
print("Standalone PDF Page Count:", len(doc_p))
