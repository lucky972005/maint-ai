import docx

doc = docx.Document(r"C:\Users\hetjo\Downloads\report3.docx")

print("Document default styles:")
for s in doc.styles:
    if s.name in ['Normal', 'List Paragraph', 'Table Grid', 'Heading 1', 'Heading 2']:
        font = s.font if hasattr(s, 'font') else None
        print(f"Style: {s.name}, font: {font.name if font else None}, size: {font.size.pt if font and font.size else None}")

print("\nSample paragraphs font info:")
for i in [0, 1, 6, 7, 11, 26, 36, 37]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        runs_info = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold) for r in p.runs[:3]]
        print(f"P[{i}]: text='{p.text[:30]}', style={p.style.name}, runs={runs_info}")

print("\nSample table cell formatting:")
t = doc.tables[0]
for r_i, r in enumerate(t.rows[:2]):
    for c_i, c in enumerate(r.cells):
        p = c.paragraphs[0]
        runs_info = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.font.bold) for r in p.runs[:2]]
        print(f"T0 r{r_i} c{c_i}: text='{c.text.strip()}', runs={runs_info}")
