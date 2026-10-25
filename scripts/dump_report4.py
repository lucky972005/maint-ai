import docx

docx_path = r"C:\dev\pro\MaintAI_Reporting_4_Het_Joshi.docx"
doc = docx.Document(docx_path)

with open(r"C:\dev\pro\report4_full_dump.txt", "w", encoding="utf-8") as f:
    f.write(f"=== PARAGRAPHS ({len(doc.paragraphs)}) ===\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P[{i}]: {p.text}\n")
        for r_i, r in enumerate(p.runs):
            if "drawing" in r._r.xml:
                f.write(f"    [RUN {r_i} HAS DRAWING]\n")

    f.write(f"\n=== TABLES ({len(doc.tables)}) ===\n")
    for i, t in enumerate(doc.tables):
        f.write(f"\n--- TABLE {i} ({len(t.rows)}x{len(t.columns)}) ---\n")
        for r_i, r in enumerate(t.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"  Row {r_i}: {cells}\n")

print("Dumped report4_full_dump.txt successfully!")
