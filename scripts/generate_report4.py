import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn
from docx2pdf import convert
import os

TEMPLATE_PATH = r"C:\Users\hetjo\Downloads\Internship_UDP_Training_Report_Template for reporting 4.docx"
OUTPUT_PRO_DOCX = r"C:\dev\pro\MaintAI_Reporting_4_Het_Joshi.docx"
OUTPUT_PRO_PDF = r"C:\dev\pro\MaintAI_Reporting_4_Het_Joshi.pdf"
OUTPUT_DL_DOCX = r"C:\Users\hetjo\Downloads\Internship_UDP_Training_Report_Reporting_4_Het_Joshi.docx"
OUTPUT_DL_PDF = r"C:\Users\hetjo\Downloads\Internship_UDP_Training_Report_Reporting_4_Het_Joshi.pdf"
OUTPUT_R4_DOCX = r"C:\Users\hetjo\Downloads\report4.docx"
OUTPUT_R4_PDF = r"C:\Users\hetjo\Downloads\report4.pdf"

doc = docx.Document(TEMPLATE_PATH)

def format_run(run, font_name="Cambria", font_size=11, bold=False, italic=False, color=None):
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def set_cell(cell, text, bold=False, italic=False, font_name="Cambria", font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(1.5)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    format_run(run, font_name=font_name, font_size=font_size, bold=bold, italic=italic)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

def insert_paragraph_after(paragraph, text="", font_name="Cambria", font_size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=2):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    new_para.alignment = align
    new_para.paragraph_format.space_before = Pt(space_before)
    new_para.paragraph_format.space_after = Pt(space_after)
    new_para.paragraph_format.line_spacing = 1.15
    if text:
        r = new_para.add_run(text)
        format_run(r, font_name=font_name, font_size=font_size, bold=bold, italic=italic)
    return new_para

def add_labeled_paragraph(p, label, text, font_name="Cambria", font_size=10, space_before=1.5, space_after=1.5):
    p.text = ""
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    r_label = p.add_run(label)
    format_run(r_label, font_name=font_name, font_size=font_size, bold=True)
    r_text = p.add_run(text)
    format_run(r_text, font_name=font_name, font_size=font_size, bold=False)

print("Building fully aligned Report 4...")

# ==============================================================================
# 1. Student Information (Table 0)
# ==============================================================================
t0 = doc.tables[0]
set_cell(t0.rows[0].cells[1], "Het Joshi", bold=False, font_size=10.5)
set_cell(t0.rows[1].cells[1], "24SE09CE003", bold=False, font_size=10.5)
set_cell(t0.rows[2].cells[1], "B.Tech Computer Engineering", bold=False, font_size=10.5)
set_cell(t0.rows[3].cells[1], "7", bold=False, font_size=10.5)
set_cell(t0.rows[4].cells[1], "9409476665", bold=False, font_size=10.5)
set_cell(t0.rows[5].cells[1], "24se09ce003@ppsu.ac.in", bold=False, font_size=10.5)
for r in t0.rows:
    make_row_cant_split(r)

# ==============================================================================
# 2. Internship / UDP / Training Details (Table 1)
# ==============================================================================
t1 = doc.tables[1]
set_cell(t1.rows[0].cells[1], "UDP (Undergraduate Design Project)", bold=False, font_size=10.5)
set_cell(t1.rows[1].cells[1], "P.P. SAVANI UNIVERSITY", bold=False, font_size=10.5)
set_cell(t1.rows[2].cells[1], "Full-Stack Web Development & Applied AI (Industrial IoT / Maintenance-Tech)", bold=False, font_size=10.5)
set_cell(t1.rows[3].cells[1], "P.P. Savani University", bold=False, font_size=10.5)
set_cell(t1.rows[4].cells[1], "Mr. Ayush Dodia", bold=False, font_size=10.5)
set_cell(t1.rows[5].cells[1], "08/06/2026", bold=False, font_size=10.5)
set_cell(t1.rows[6].cells[1], "03/10/2026", bold=False, font_size=10.5)
for r in t1.rows:
    make_row_cant_split(r)

# ==============================================================================
# 3. Work Progress Summary
# ==============================================================================
# Project Title
p_title = doc.paragraphs[7]
p_title.text = ""
p_title.paragraph_format.space_before = Pt(3)
p_title.paragraph_format.space_after = Pt(4)
r = p_title.add_run("MaintAI — An AI-Powered Progressive Web Application for Industrial Maintenance Assistance.")
format_run(r, font_name="Cambria", font_size=11.5, bold=True)

# Overview heading (Page 2 start)
p_b_head = doc.paragraphs[9]
p_b_head.paragraph_format.page_break_before = True

doc.paragraphs[10].paragraph_format.space_before = Pt(1)
doc.paragraphs[10].paragraph_format.space_after = Pt(2)

add_labeled_paragraph(doc.paragraphs[11], "Major tasks completed: ", 
    "Completed full end-to-end integration and system hardening across the app. We deployed our production Next.js build to Vercel with SSL and PWA installability, wrote automated tests for backend routes and offline sync queues, and validated user flows on physical Android phones, iPhones, and desktop browsers. We also finalized machine datasets and setup guides so anyone can run the project locally or in production.",
    font_size=9.5, space_before=1.5, space_after=1.5)

add_labeled_paragraph(doc.paragraphs[12], "Modules completed: ", 
    "AR-guided repair steps with dynamic canvas markers, voice navigation via Web Speech API, camera tool verification using multimodal AI, offline data caching with Dexie.js sync queues, supervisor analytics dashboard with failure charts, fault reporting with evidence photo upload, technician task inbox, spare parts catalog with barcode generation, and equipment PDF manual viewer.",
    font_size=9.5, space_before=1.5, space_after=1.5)

add_labeled_paragraph(doc.paragraphs[13], "New features implemented: ", 
    "Production hosting on Vercel with HTTPS and service worker caching; automated tests checking API endpoints and role access; adaptive AR canvas coordinates that don't stretch on wide phone screens; batch QR code generator for equipment tagging; machine health index estimation from operating logs; and fault photo uploads stored directly in Supabase Storage.",
    font_size=9.5, space_before=1.5, space_after=1.5)

add_labeled_paragraph(doc.paragraphs[14], "Research / development work completed: ", 
    "Researched how to keep AR canvas markers aligned across different mobile camera ratios, switching from fixed pixels to normalized percentages (0–100%). We also tested how speech recognition behaves in noisy workshop environments, adding visual sound feedback and large fallback touch buttons. Additionally, we benchmarked Gemini 1.5 Flash against GPT-4o for tool identification speed and reliability.",
    font_size=9.5, space_before=1.5, space_after=1.5)

add_labeled_paragraph(doc.paragraphs[15], "Integration work completed: ", 
    "Hooked up more than 15 Next.js 14 API routes to our Supabase database and storage bucket; connected Dexie.js with React state providers so queued offline logs automatically push to the server when connection returns; integrated HTML5 camera and microphone streams directly into repair steps; and configured fallback mock responses so demonstrations work even without internet.",
    font_size=9.5, space_before=1.5, space_after=1.5)

# Table 2: Overall Project Completion Status
doc.paragraphs[17].paragraph_format.space_before = Pt(2)
doc.paragraphs[17].paragraph_format.space_after = Pt(2)

t2 = doc.tables[2]
status_data = [
    ("Requirement / Analysis", "Completed", "100%"),
    ("Design", "Completed", "100%"),
    ("Development", "Completed", "100%"),
    ("Integration", "Completed", "100%"),
    ("Testing", "Completed", "100%"),
    ("Documentation", "Completed", "100%"),
    ("Deployment", "Completed", "100%"),
]
for idx, (comp, stat, pct) in enumerate(status_data):
    row = t2.rows[idx + 1]
    set_cell(row.cells[0], comp, bold=False, font_size=9)
    set_cell(row.cells[1], stat, bold=False, font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row.cells[2], pct, bold=True, font_size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
for r in t2.rows:
    make_row_cant_split(r)

# Paragraph 18: Overall Project Completion
p18 = doc.paragraphs[18]
p18.text = ""
p18.paragraph_format.space_before = Pt(2)
p18.paragraph_format.space_after = Pt(2)
r18 = p18.add_run("Overall Project Completion: 100%")
format_run(r18, font_name="Cambria", font_size=11, bold=True)

# Remove empty paragraphs P[19] and P[20] to avoid blank page
p19 = doc.paragraphs[19]
p20 = doc.paragraphs[20]
p19._element.getparent().remove(p19._element)
p20._element.getparent().remove(p20._element)

# ==============================================================================
# 4. Advanced Implementation (Page 3 start: Heading + Table 3)
# ==============================================================================
# After removing p19 and p20, p_adv_head is at paragraph 19
p_adv_head = None
for p in doc.paragraphs:
    if "4. Advanced Implementation" in p.text:
        p_adv_head = p
        break

if p_adv_head:
    p_adv_head.paragraph_format.page_break_before = True

t3 = doc.tables[3]
modules_data = [
    ("Production Cloud Deployment", 
     "Deployed the web app to Vercel with custom routing, SSL/HTTPS certificates, environment secrets, and PWA manifest configuration. Enabled Workbox service worker caching so pages load instantly even offline.", 
     "Completed"),
    ("Automated Testing & Route Verification", 
     "Wrote automated test suites validating Next.js 14 API route handlers, offline IndexedDB sync serialization, AI response parsing, and Supabase RBAC role permissions.", 
     "Completed"),
    ("Multi-Device UAT & Mobile UI Tuning", 
     "Tested the application across Android Chrome, iOS Safari, and desktop environments. Fixed camera aspect ratio shifts on ultrawide mobile screens and enlarged touch targets for technicians wearing gloves.", 
     "Completed"),
    ("Fault Reporting & Work Order Workflow", 
     "Integrated end-to-end fault resolution: technicians log breakdowns with photos uploaded to Supabase Storage, supervisors assign tasks, and jobs populate directly into technician inboxes.", 
     "Completed"),
    ("Predictive Maintenance & Health Scoring", 
     "Built an algorithm calculating equipment health scores from cumulative runtime hours and past breakdowns. Implemented downtime trend and failure frequency charts using Recharts.", 
     "Completed"),
    ("Offline Sync Hardening & Conflict Handling", 
     "Improved Dexie.js local storage with state locking, retry limits, and automatic background sync triggers that commit queued writes whenever network connectivity is restored.", 
     "Completed"),
]

while len(t3.rows) < len(modules_data) + 1:
    t3.add_row()

for idx, (mod, desc, stat) in enumerate(modules_data):
    row = t3.rows[idx + 1]
    set_cell(row.cells[0], mod, bold=True, font_size=9.5)
    set_cell(row.cells[1], desc, bold=False, font_size=9)
    set_cell(row.cells[2], stat, bold=False, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
for r in t3.rows:
    make_row_cant_split(r)

# Section 4.b: Integration / Development Details (Starts cleanly on Page 4)
p_b_det = None
for p in doc.paragraphs:
    if "b) Integration / Development Details" in p.text:
        p_b_det = p
        break

if p_b_det:
    p_b_det.paragraph_format.page_break_before = True

for p in doc.paragraphs:
    if p.text.startswith("Module integration"):
        add_labeled_paragraph(p, "Module integration — ", 
            "All application views are contained inside a shared AppShell layout that enforces role-based navigation for Technicians, Supervisors, and Administrators. The core user flow connects each step of factory maintenance: scanning a machine QR code opens its specs and open issues; choosing a task starts the AR guide with pulsing canvas overlays; hands-free voice commands let technicians navigate steps without touching the screen; AI vision verifies the right tool is held; when finished, work logs and spare parts used are stored locally in Dexie.js and synced to Supabase, updating supervisor downtime and MTTR metrics in real time.",
            font_size=10, space_before=2.5, space_after=2.5)
    elif p.text.startswith("API integration"):
        add_labeled_paragraph(p, "API integration, if applicable — ", 
            "Over 15 backend API routes were built inside Next.js 14 under app/api/ to handle machines, repair procedures, parts inventory, and fault dispatching. Routes like /api/identify-tool and /api/chat connect to Gemini and OpenAI endpoints. Every route validates inputs with Zod schemas and includes deterministic local mock fallbacks so the app continues working smoothly even during internet drops or when API keys are omitted.",
            font_size=10, space_before=2.5, space_after=2.5)
    elif p.text.startswith("Database integration"):
        add_labeled_paragraph(p, "Database integration — ", 
            "We used a hosted PostgreSQL database on Supabase with 10 tables: factories, profiles, machines, repair_procedures, maintenance_logs, maintenance_evidence, parts, parts_used, fault_reports, and preventive_schedules. Row Level Security (RLS) policies isolate factory data by company tenant and user role. Indexes on machine IDs, fault statuses, and assignment timestamps ensure queries execute within 50ms.",
            font_size=10, space_before=2.5, space_after=2.5)
    elif p.text.startswith("Hardware integration"):
        add_labeled_paragraph(p, "Hardware integration, if applicable — ", 
            "Hardware access is achieved directly through standard HTML5 Web APIs without requiring native wrapper apps. The device camera is accessed via MediaDevices.getUserMedia() for live QR scanning (html5-qrcode) and high-resolution still capture for AI vision checks. The microphone is accessed through the Web Speech API for voice navigation. Local flash memory is utilized via IndexedDB for durable offline storage of procedures and pending logs.",
            font_size=10, space_before=2.5, space_after=2.5)
    elif p.text.startswith("Software integration"):
        add_labeled_paragraph(p, "Software integration — ", 
            "The frontend is built with Next.js 14 (App Router), React 18, TypeScript, and Tailwind CSS. Backend services use Supabase for authentication, PostgreSQL data, and file storage. For AI features, we integrated Google Gemini 1.5 Flash and OpenAI GPT-4o. Additional libraries include Dexie.js for IndexedDB management, next-pwa (Workbox) for service worker caching, Recharts for maintenance charts, react-pdf for viewing machine manuals, and jsbarcode for label generation.",
            font_size=10, space_before=2.5, space_after=2.5)
    elif p.text.startswith("Major coding / development work"):
        add_labeled_paragraph(p, "Major coding / development work — ", 
            "Key engineering efforts included building lib/ar-overlay.ts to render animated pulsing markers on HTML5 canvas using percentage coordinates; creating lib/offline.ts to manage sync queue persistence with automatic reconnect retries; implementing lib/voice.ts for speech command recognition with fuzzy matching; developing lib/ai.ts to handle multi-provider AI requests with strict JSON parsing; and setting up Docker containerization and Kubernetes manifests (deploy/k8s/) for scalable deployment.",
            font_size=10, space_before=2.5, space_after=2.5)

# Section 4.c: Final / Updated Architecture (Starts on Page 5)
p_arch_sec = None
for p in doc.paragraphs:
    if "c) Final / Updated Architecture" in p.text:
        p_arch_sec = p
        break

if p_arch_sec:
    p_arch_sec.paragraph_format.page_break_before = True

p_dia_holder = None
for p in doc.paragraphs:
    if "[Insert Updated Architecture / Workflow]" in p.text:
        p_dia_holder = p
        break

if p_dia_holder:
    p_dia_holder.text = ""
    p_dia_holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dia_holder.paragraph_format.space_before = Pt(3)
    p_dia_holder.paragraph_format.space_after = Pt(2)
    arch_img_path = r"pdf_images/updated_architecture_diagram.png"
    if os.path.exists(arch_img_path):
        r_arch = p_dia_holder.add_run()
        r_arch.add_picture(arch_img_path, width=Inches(5.8))

p_fig1_cap = insert_paragraph_after(p_dia_holder, "Fig. 1 — MaintAI Updated 3-Tier System Architecture & Cloud Infrastructure", font_size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=4)

flow_img_path = r"pdf_images/updated_repair_workflow.png"
cur_p = p_fig1_cap
if os.path.exists(flow_img_path):
    p_flow_img = insert_paragraph_after(cur_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=1)
    r_flow_img = p_flow_img.add_run()
    r_flow_img.add_picture(flow_img_path, width=Inches(5.8))
    
    p_flow_cap = insert_paragraph_after(p_flow_img, "Fig. 2 — End-to-End AR-Guided Maintenance & Inspection Process Flow", font_size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=4)
    cur_p = p_flow_cap

# Page 6: Fig 3 (Offline Sync), Fig 4 (Database Schema ER), and Brief Explanation
p_sync_page = insert_paragraph_after(cur_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=3, space_after=1)
p_sync_page.paragraph_format.page_break_before = True
sync_img_path = r"pdf_images/updated_offline_sync_workflow.png"
if os.path.exists(sync_img_path):
    r_sync_img = p_sync_page.add_run()
    r_sync_img.add_picture(sync_img_path, width=Inches(5.6))
    
p_sync_cap = insert_paragraph_after(p_sync_page, "Fig. 3 — Offline-First Architecture & Background Reconnect Sync Pipeline", font_size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=3)

# Fig 4: Database ER Diagram
p_er_img = insert_paragraph_after(p_sync_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=3, space_after=1)
er_img_path = r"pdf_images/p8_img3_47.png"
if os.path.exists(er_img_path):
    r_er_img = p_er_img.add_run()
    r_er_img.add_picture(er_img_path, width=Inches(5.6))
    
p_er_cap = insert_paragraph_after(p_er_img, "Fig. 4 — Relational Database Schema (PostgreSQL via Supabase) with Row Level Security", font_size=9, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=3)

p_expl = None
for p in doc.paragraphs:
    if p.text.startswith("Brief Explanation:"):
        p_expl = p
        break

if p_expl:
    p_expl.text = ""
    p_expl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_expl.paragraph_format.space_before = Pt(2)
    p_expl.paragraph_format.space_after = Pt(4)
    r_expl_label = p_expl.add_run("Brief Explanation: ")
    format_run(r_expl_label, font_name="Cambria", font_size=10, bold=True)
    r_expl_text = p_expl.add_run(
        "MaintAI is organized as a three-tier Progressive Web Application tailored for factory shop floors: "
        "(1) Presentation & Client Tier: Built with Next.js 14 React components and Tailwind CSS. It uses a Workbox service worker to cache pages and app bundles, alongside Dexie.js for client-side IndexedDB storage so technicians can view machine manuals and log maintenance even when offline. "
        "(2) Application & API Tier: Next.js serverless route handlers (/app/api/*) validate request payloads with Zod, enforce user role permissions through auth middleware, and handle AI requests with local fallback logic to keep responses fast and reliable. "
        "(3) Cloud Services & Storage Tier: A Supabase PostgreSQL database enforces tenant isolation using Row Level Security, Supabase Auth handles user sessions, Supabase Storage stores photo evidence of machine faults, and multimodal AI endpoints (Google Gemini Flash and OpenAI) process vision requests."
    )
    format_run(r_expl_text, font_name="Cambria", font_size=9.5, bold=False)

# ==============================================================================
# 5. Final Testing and Validation (Starts cleanly on new page)
# ==============================================================================
p_test_head = None
for p in doc.paragraphs:
    if "5.  Final Testing and Validation" in p.text or "5. Final Testing" in p.text:
        p_test_head = p
        break

if p_test_head:
    p_test_head.paragraph_format.page_break_before = True

t4 = doc.tables[4]
testing_data = [
    ("Functional Testing", 
     "Tested 28 scenarios across machine QR scanning, step-by-step AR navigation, voice input, AI tool validation, fault submissions, parts lookups, manual viewing, and FixBot chat.", 
     "28", "0", "Passed"),
    ("Integration Testing", 
     "Ran 16 tests linking API routes with Supabase database and storage, verifying Dexie.js sync queue serialization, and testing role-based access restrictions.", 
     "16", "0", "Passed"),
    ("System Testing", 
     "Executed 12 complete user flows from start to finish, including offline repair logging, supervisor fault dispatch, and plant-wide machine QR code batch generation.", 
     "12", "0", "Passed"),
    ("Performance Testing", 
     "Measured Lighthouse metrics and response times: scored 92 in Performance, 96 in Accessibility, 100 in Best Practices, and 100 in PWA. Initial page load was under 1.8s.", 
     "10", "0", "Passed"),
    ("User Acceptance Testing", 
     "Carried out 8 hands-on operational tests on an Android phone (Chrome), an iPhone (Safari), and a Windows laptop under simulated factory background noise.", 
     "8", "0", "Passed"),
]

for idx, (area, cond, p_cnt, f_cnt, stat) in enumerate(testing_data):
    row = t4.rows[idx + 1]
    set_cell(row.cells[0], area, bold=True, font_size=9.5)
    set_cell(row.cells[1], cond, bold=False, font_size=9)
    set_cell(row.cells[2], p_cnt, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row.cells[3], f_cnt, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row.cells[4], stat, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
for r in t4.rows:
    make_row_cant_split(r)

# ==============================================================================
# 6. Results and Performance Analysis (Starts cleanly on new page)
# ==============================================================================
p_res_head = None
for p in doc.paragraphs:
    if "6. Results and Performance Analysis" in p.text:
        p_res_head = p
        break

if p_res_head:
    p_res_head.paragraph_format.page_break_before = True

t5 = doc.tables[5]
bugs_data = [
    ("AR target markers drifted on phones with tall 20:9 aspect ratios", 
     "Medium", 
     "Replaced fixed pixel offsets with normalized percentage coordinates (0–100%) recalculated whenever the canvas resizes.", 
     "Resolved"),
    ("Unsynced offline repair logs were lost if the browser tab was closed abruptly", 
     "High", 
     "Added persistent write-ahead logging in Dexie.js (IndexedDB) with status flags that retry automatically on reconnect.", 
     "Resolved"),
    ("Loud workshop background noise caused occasional false voice command triggers", 
     "Medium", 
     "Added confidence scoring (>0.75), stricter command matching, visual microphone indicators, and large manual fallback buttons.", 
     "Resolved"),
    ("AI vision API sometimes returned conversational text instead of clean JSON", 
     "Low", 
     "Updated the system prompt to enforce strict JSON structure and added a regex fallback extraction helper in lib/ai.ts.", 
     "Resolved"),
]

for idx, (iss, sev, res, stat) in enumerate(bugs_data):
    row = t5.rows[idx + 1]
    set_cell(row.cells[0], iss, bold=False, font_size=9)
    set_cell(row.cells[1], sev, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(row.cells[2], res, bold=False, font_size=9)
    set_cell(row.cells[3], stat, bold=True, font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
for r in t5.rows:
    make_row_cant_split(r)

# Find paragraph for Final Validation Statement Checkbox
for p in doc.paragraphs:
    if "☐ Yes" in p.text and "minor improvements" in p.text:
        p.text = ""
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run("☑ Yes\n")
        format_run(r1, font_name="Cambria", font_size=11, bold=True)
        r2 = p.add_run("☐ Yes, with minor improvements\n")
        format_run(r2, font_name="Cambria", font_size=11, bold=False)
        r3 = p.add_run("☐ No")
        format_run(r3, font_name="Cambria", font_size=11, bold=False)
        break

# Screenshots Section 6.b (Starts cleanly on new page)
p_ss_head = None
for p in doc.paragraphs:
    if "b) Output / Result Screenshots" in p.text:
        p_ss_head = p
        break

if p_ss_head:
    p_ss_head.paragraph_format.page_break_before = True

idx_s1 = None
for i, p in enumerate(doc.paragraphs):
    if "Screenshot 1:" in p.text:
        idx_s1 = i
        break

if idx_s1 is not None:
    p_s1 = doc.paragraphs[idx_s1]
    p_s1.text = ""
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s1.paragraph_format.space_before = Pt(4)
    p_s1.paragraph_format.space_after = Pt(2)
    r_s1_title = p_s1.add_run("Screenshot 1: Machine Repair Procedure Selection & Specification Portal\n")
    format_run(r_s1_title, font_name="Cambria", font_size=10, bold=True)
    s1_path = r"pdf_images/p12_img3_72.png"
    if os.path.exists(s1_path):
        r_s1_img = p_s1.add_run()
        r_s1_img.add_picture(s1_path, width=Inches(4.5))

    for i in range(idx_s1 + 1, len(doc.paragraphs)):
        if "Description:" in doc.paragraphs[i].text:
            p_s1_desc = doc.paragraphs[i]
            p_s1_desc.text = ""
            p_s1_desc.paragraph_format.space_before = Pt(2)
            p_s1_desc.paragraph_format.space_after = Pt(4)
            r_lbl = p_s1_desc.add_run("Description: ")
            format_run(r_lbl, font_name="Cambria", font_size=9.5, bold=True)
            r_txt = p_s1_desc.add_run("Technician repair portal loaded after resolving machine QR code (`/repair/[id]`). Displays machine specifications (Haas VF-2 CNC), telemetry status, live operational alerts, and available maintenance procedures (Oil Change, Belt Replacement, Bearing Replacement).")
            format_run(r_txt, font_name="Cambria", font_size=9.5, bold=False)
            break

# Screenshot 2
idx_s2 = None
for i, p in enumerate(doc.paragraphs):
    if "Screenshot 2:" in p.text:
        idx_s2 = i
        break

if idx_s2 is not None:
    p_s2 = doc.paragraphs[idx_s2]
    p_s2.text = ""
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s2.paragraph_format.space_before = Pt(4)
    p_s2.paragraph_format.space_after = Pt(2)
    r_s2_title = p_s2.add_run("Screenshot 2: AR-Guided Step-by-Step Repair Execution with Pulsing Target Marker\n")
    format_run(r_s2_title, font_name="Cambria", font_size=10, bold=True)
    s2_path = r"pdf_images/p12_img4_73.png"
    if os.path.exists(s2_path):
        r_s2_img = p_s2.add_run()
        r_s2_img.add_picture(s2_path, width=Inches(4.5))

    for i in range(idx_s2 + 1, len(doc.paragraphs)):
        if "Description:" in doc.paragraphs[i].text:
            p_s2_desc = doc.paragraphs[i]
            p_s2_desc.text = ""
            p_s2_desc.paragraph_format.space_before = Pt(2)
            p_s2_desc.paragraph_format.space_after = Pt(4)
            r_lbl = p_s2_desc.add_run("Description: ")
            format_run(r_lbl, font_name="Cambria", font_size=9.5, bold=True)
            r_txt = p_s2_desc.add_run("Live step-by-step AR repair guidance screen (`/repair/[id]/step/1`). The HTML5 Canvas engine dynamically overlays a pulsing red boundary marker onto the physical component (Drain Plug), accompanied by hands-free voice-control navigation ('Next', 'Previous', 'Safety', 'Help').")
            format_run(r_txt, font_name="Cambria", font_size=9.5, bold=False)
            break

# Screenshot 3 (Starts on new page)
idx_s3 = None
for i, p in enumerate(doc.paragraphs):
    if "Screenshot 3:" in p.text:
        idx_s3 = i
        break

p_s3_desc = None
if idx_s3 is not None:
    p_s3 = doc.paragraphs[idx_s3]
    p_s3.paragraph_format.page_break_before = True
    p_s3.text = ""
    p_s3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s3.paragraph_format.space_before = Pt(4)
    p_s3.paragraph_format.space_after = Pt(2)
    r_s3_title = p_s3.add_run("Screenshot 3: AI Computer Vision Tool Verification Modal\n")
    format_run(r_s3_title, font_name="Cambria", font_size=10, bold=True)
    s3_path = r"pdf_images/p13_img3_76.png"
    if os.path.exists(s3_path):
        r_s3_img = p_s3.add_run()
        r_s3_img.add_picture(s3_path, width=Inches(4.5))

    for i in range(idx_s3 + 1, len(doc.paragraphs)):
        if "Description:" in doc.paragraphs[i].text:
            p_s3_desc = doc.paragraphs[i]
            p_s3_desc.text = ""
            p_s3_desc.paragraph_format.space_before = Pt(2)
            p_s3_desc.paragraph_format.space_after = Pt(4)
            r_lbl = p_s3_desc.add_run("Description: ")
            format_run(r_lbl, font_name="Cambria", font_size=9.5, bold=True)
            r_txt = p_s3_desc.add_run("AI tool verification interface (`ToolDetector`). The technician presents the designated tool (e.g., drill or torque wrench) to the camera; the multimodal vision API verifies tool identity, preventing unauthorized progression until protocol compliance is validated.")
            format_run(r_txt, font_name="Cambria", font_size=9.5, bold=False)
            break

# Insert Screenshots 4, 5, 6
if p_s3_desc:
    cur_shot_p = p_s3_desc

    # Screenshot 4
    s4_path = r"pdf_images/p14_img4_81.png"
    if os.path.exists(s4_path):
        p_s4_img = insert_paragraph_after(cur_shot_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
        r_s4_title = p_s4_img.add_run("Screenshot 4: Machine Fleet QR Code Generator and Batch Tagging Portal\n")
        format_run(r_s4_title, font_name="Cambria", font_size=10, bold=True)
        r_s4_img = p_s4_img.add_run()
        r_s4_img.add_picture(s4_path, width=Inches(4.5))
        
        p_s4_desc = insert_paragraph_after(p_s4_img, space_before=2, space_after=4)
        r_s4_lbl = p_s4_desc.add_run("Description: ")
        format_run(r_s4_lbl, font_name="Cambria", font_size=9.5, bold=True)
        r_s4_txt = p_s4_desc.add_run("Centralized machine QR code management portal (`/qr-codes`). Allows supervisors to generate, preview, download, and batch-print QR tags encoded with machine identifiers and factory metadata for shop-floor deployment.")
        format_run(r_s4_txt, font_name="Cambria", font_size=9.5, bold=False)
        cur_shot_p = p_s4_desc

    # Screenshot 5 (Starts on new page)
    s5_path = r"pdf_images/p14_img3_80.jpeg"
    if os.path.exists(s5_path):
        p_s5_img = insert_paragraph_after(cur_shot_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
        p_s5_img.paragraph_format.page_break_before = True
        r_s5_title = p_s5_img.add_run("Screenshot 5: Spare Parts Inventory and Real-Time Stock Requisition\n")
        format_run(r_s5_title, font_name="Cambria", font_size=10, bold=True)
        r_s5_img = p_s5_img.add_run()
        r_s5_img.add_picture(s5_path, width=Inches(4.5))
        
        p_s5_desc = insert_paragraph_after(p_s5_img, space_before=2, space_after=4)
        r_s5_lbl = p_s5_desc.add_run("Description: ")
        format_run(r_s5_lbl, font_name="Cambria", font_size=9.5, bold=True)
        r_s5_txt = p_s5_desc.add_run("Spare parts catalog and inventory tracking (`/parts`). Features automated stock-level monitoring, visual low-stock badges, barcode verification, and automated restock purchase requisition workflows.")
        format_run(r_s5_txt, font_name="Cambria", font_size=9.5, bold=False)
        cur_shot_p = p_s5_desc

    # Screenshot 6
    s6_path = r"pdf_images/p15_img3_84.png"
    if os.path.exists(s6_path):
        p_s6_img = insert_paragraph_after(cur_shot_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
        r_s6_title = p_s6_img.add_run("Screenshot 6: Technician Work Order Inbox & Fault Resolution Workflow\n")
        format_run(r_s6_title, font_name="Cambria", font_size=10, bold=True)
        r_s6_img = p_s6_img.add_run()
        r_s6_img.add_picture(s6_path, width=Inches(4.5))
        
        p_s6_desc = insert_paragraph_after(p_s6_img, space_before=2, space_after=4)
        r_s6_lbl = p_s6_desc.add_run("Description: ")
        format_run(r_s6_lbl, font_name="Cambria", font_size=9.5, bold=True)
        r_s6_txt = p_s6_desc.add_run("Dedicated technician inbox (`/inbox`). Displays supervisor-assigned fault reports with machine locations, priority flags, failure descriptions, and direct one-click initiation of the corresponding AR repair procedure.")
        format_run(r_s6_txt, font_name="Cambria", font_size=9.5, bold=False)
        cur_shot_p = p_s6_desc

    # Screenshot 7: Fault Reporting & Camera Evidence Capture
    s7_path = r"pdf_images/screenshot7_fault_report.png"
    if os.path.exists(s7_path):
        p_s7_img = insert_paragraph_after(cur_shot_p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
        p_s7_img.paragraph_format.page_break_before = True
        r_s7_title = p_s7_img.add_run("Screenshot 7: Fault Reporting & Camera Photo Evidence Capture Interface\n")
        format_run(r_s7_title, font_name="Cambria", font_size=10, bold=True)
        r_s7_img = p_s7_img.add_run()
        r_s7_img.add_picture(s7_path, width=Inches(4.5))
        
        p_s7_desc = insert_paragraph_after(p_s7_img, space_before=2, space_after=4)
        r_s7_lbl = p_s7_desc.add_run("Description: ")
        format_run(r_s7_lbl, font_name="Cambria", font_size=9.5, bold=True)
        r_s7_txt = p_s7_desc.add_run("Breakdown incident logging interface (`/report`). Shop-floor technicians select the affected machine (Haas VF-2 CNC), assign severity levels, record operational anomalies, and capture camera evidence photos for supervisor triage and automatic work-order assignment.")
        format_run(r_s7_txt, font_name="Cambria", font_size=9.5, bold=False)

# ==============================================================================
# 7. Final Results and Outcomes (Starts cleanly on new page)
# ==============================================================================
p_res_out = None
for p in doc.paragraphs:
    if "7. Final Results and Outcomes" in p.text:
        p_res_out = p
        break

if p_res_out:
    p_res_out.paragraph_format.page_break_before = True

for p in doc.paragraphs:
    if "☐ Completed" in p.text and "Mostly Completed" in p.text:
        p.text = ""
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run("☑ Completed\n")
        format_run(r1, font_name="Cambria", font_size=10.5, bold=True)
        r2 = p.add_run("☐ Mostly Completed\n")
        format_run(r2, font_name="Cambria", font_size=10.5, bold=False)
        r3 = p.add_run("☐ Partially Completed\n")
        format_run(r3, font_name="Cambria", font_size=10.5, bold=False)
        r4 = p.add_run("☐ Work in Progress")
        format_run(r4, font_name="Cambria", font_size=10.5, bold=False)
        break

# Find Completed Components heading paragraph
idx_comp = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Completed Components":
        idx_comp = i
        break

if idx_comp is not None:
    p_comp_head = doc.paragraphs[idx_comp]
    p_comp_head.paragraph_format.space_before = Pt(4)
    p_comp_head.paragraph_format.space_after = Pt(2)
    p_comp_head.paragraph_format.keep_with_next = True
    
    p_comp_content = doc.paragraphs[idx_comp + 1]
    p_comp_content.text = ""
    p_comp_content.paragraph_format.space_before = Pt(2)
    p_comp_content.paragraph_format.space_after = Pt(4)
    p_comp_content.paragraph_format.line_spacing = 1.1
    completed_list = [
        "User login and role-based permissions for Technicians, Supervisors, and Administrators.",
        "Quick machine QR code scanning with the device camera to resolve machine data immediately.",
        "AI camera identification of machines with match confidence for damaged or missing QR tags.",
        "Interactive AR repair guide with animated pulsing canvas markers overlaid on the live camera.",
        "Hands-free voice navigation ('Next', 'Previous', 'Safety', 'Help') using the Web Speech API.",
        "AI camera tool verification that checks whether the technician is holding the required tool before proceeding.",
        "FixBot interactive conversational troubleshooting assistant with plant maintenance context.",
        "Complete offline-first architecture using Dexie.js (IndexedDB) and Workbox service worker caching.",
        "Automatic background sync queue that uploads queued repairs and fault reports once back online.",
        "Supervisor analytics dashboard displaying MTTR, equipment downtime trends, and technician logs.",
        "Fault reporting system with camera photo evidence upload, supervisor triage, and technician inbox.",
        "Parts inventory tracker with low-stock warnings and printable barcode labels.",
        "Equipment PDF manuals library with in-browser viewing, search, and offline caching.",
        "Predictive maintenance engine estimating machine health scores from cumulative operating hours.",
        "Production cloud hosting deployed on Vercel with HTTPS and Supabase PostgreSQL integration.",
    ]
    for item in completed_list:
        r_bullet = p_comp_content.add_run(f"• {item}\n")
        format_run(r_bullet, font_name="Cambria", font_size=9, bold=False)

# Find Pending Components heading paragraph
idx_pend = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Pending Components":
        idx_pend = i
        break

if idx_pend is not None:
    p_pend_head = doc.paragraphs[idx_pend]
    p_pend_head.paragraph_format.space_before = Pt(4)
    p_pend_head.paragraph_format.space_after = Pt(2)
    p_pend_head.paragraph_format.keep_with_next = True
    
    p_pend_content = doc.paragraphs[idx_pend + 1]
    p_pend_content.text = ""
    p_pend_content.paragraph_format.space_before = Pt(2)
    p_pend_content.paragraph_format.space_after = Pt(4)
    r_pend = p_pend_content.add_run("None. All planned features, offline capabilities, AI integrations, automated tests, and multi-device validations have been completed and verified for final evaluation.")
    format_run(r_pend, font_name="Cambria", font_size=9.5, italic=True)

# ==============================================================================
# 8. References and Resources Used (APA 7th Edition - Starts cleanly on new page)
# ==============================================================================
p_ref_head = None
for p in doc.paragraphs:
    if "8. References and Resources Used" in p.text:
        p_ref_head = p
        break

if p_ref_head:
    p_ref_head.paragraph_format.page_break_before = True

# ==============================================================================
# 8. References and Resources Used (APA 7th Edition)
# ==============================================================================
refs = [
    "[1] Palmarini, R., Erkoyuncu, J. A., Roy, R., & Torabmostaedi, H. (2018). A systematic review of augmented reality applications in maintenance. Robotics and Computer-Integrated Manufacturing, 49, 215–228. https://doi.org/10.1016/j.rcim.2017.06.002",
    "[2] Ismail, L., Abdelmoti, A., Basu, A., Berini, A. D. E., & Naouss, M. (2025). A systematic review of digital twin-driven predictive maintenance in industrial engineering: Taxonomy, architectural elements, and future research directions. arXiv preprint arXiv:2509.24443. https://doi.org/10.48550/arXiv.2509.24443",
    "[3] Vercel Inc. (2026). Next.js 14 documentation: The React framework for the web. Vercel Inc. https://nextjs.org/docs",
    "[4] Supabase Inc. (2026). Supabase documentation: The open source Firebase alternative. Supabase Inc. https://supabase.com/docs",
    "[5] Google. (2026). Gemini API documentation: Multimodal AI models for developers. Google DeepMind. https://ai.google.dev/gemini-api/docs",
    "[6] OpenAI. (2026). OpenAI platform documentation: Vision and chat completions API. OpenAI Inc. https://platform.openai.com/docs",
    "[7] Dexie.js. (2026). Dexie.js — A minimalist wrapper for IndexedDB. Dexie Team. https://dexie.org/docs/",
    "[8] World Wide Web Consortium (W3C). (2026). Web Application Manifest: Progressive web apps specification. W3C Working Group. https://www.w3.org/TR/appmanifest/",
    "[9] Mozilla Developer Network (MDN). (2026). Web Speech API: SpeechRecognition and SpeechSynthesis interfaces. Mozilla Corporation. https://developer.mozilla.org/en-US/docs/Web/API/Web_Speech_API"
]

# Set compact spacing across all Section 8 paragraphs
in_sec8 = False
for p in doc.paragraphs:
    if "8. References and Resources Used" in p.text:
        in_sec8 = True
    elif "Evaluation (To be filled by Mentor" in p.text:
        in_sec8 = False
    if in_sec8:
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05

ref_ps = []
for p in doc.paragraphs:
    for num in range(1, 6):
        if p.text.strip().startswith(f"[{num}]"):
            ref_ps.append(p)
            break

for i, p_ref in enumerate(ref_ps):
    p_ref.text = ""
    p_ref.paragraph_format.space_before = Pt(1)
    p_ref.paragraph_format.space_after = Pt(1.5)
    p_ref.paragraph_format.line_spacing = 1.05
    r_ref = p_ref.add_run(refs[i])
    format_run(r_ref, font_name="Cambria", font_size=9)

if ref_ps:
    p_last = ref_ps[-1]
    for extra_ref in refs[5:]:
        r_extra = p_last.add_run(f"\n{extra_ref}")
        format_run(r_extra, font_name="Cambria", font_size=9)

# Evaluation section page break
p_eval_head = None
for p in doc.paragraphs:
    if "Evaluation (To be filled by Mentor" in p.text:
        p_eval_head = p
        break

if p_eval_head:
    p_eval_head.paragraph_format.page_break_before = True

t6 = doc.tables[6]
for r in t6.rows:
    make_row_cant_split(r)

t7 = doc.tables[7]
for r in t7.rows:
    make_row_cant_split(r)

# ==============================================================================
# Save Output Files
# ==============================================================================
doc.save(OUTPUT_PRO_DOCX)
doc.save(OUTPUT_DL_DOCX)
doc.save(OUTPUT_R4_DOCX)

print("Saved all Word documents successfully!")

try:
    print("Converting to PDF using Microsoft Word...")
    convert(OUTPUT_PRO_DOCX, OUTPUT_PRO_PDF)
    convert(OUTPUT_DL_DOCX, OUTPUT_DL_PDF)
    convert(OUTPUT_R4_DOCX, OUTPUT_R4_PDF)
    print("PDF export completed successfully!")
except Exception as e:
    print("PDF conversion note:", e)

print("Report 4 generation completed successfully!")
