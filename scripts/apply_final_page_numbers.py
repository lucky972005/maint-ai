import os
import shutil
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client

OUTPUT_PRO_DOCX = r"C:\dev\pro\MaintAI_Final_Hard_Bound_Project_Report_Het_Joshi.docx"
OUTPUT_PRO_PDF = r"C:\dev\pro\MaintAI_Final_Hard_Bound_Project_Report_Het_Joshi.pdf"
OUTPUT_DL_DOCX = r"C:\Users\hetjo\Downloads\PPSU_Final_Hard_Bound_Project_Report_Het_Joshi.docx"
OUTPUT_DL_PDF = r"C:\Users\hetjo\Downloads\PPSU_Final_Hard_Bound_Project_Report_Het_Joshi.pdf"

print("Opening docx to inject exact discovered 64-page numbers...")
doc = docx.Document(OUTPUT_PRO_DOCX)

def format_run(run, font_name="Cambria", font_size=9.0, bold=False, italic=False):
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic

def set_cell(cell, text, bold=False, italic=False, font_name="Cambria", font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = cell.paragraphs[0]
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(2.0)
    p.paragraph_format.space_after = Pt(2.0)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    format_run(run, font_name=font_name, font_size=font_size, bold=bold, italic=italic)

def set_para_text(p, text, font_name="Cambria", font_size=9.5, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=4):
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        format_run(r, font_name=font_name, font_size=font_size, bold=bold, italic=italic)

PAGE_MAP = {
    # Preliminary Pages
    'Certificate': 'i',
    'Declaration': 'ii',
    'Acknowledgement': 'iii',
    'Abstract': 'iv',
    'Table of Contents': 'v',
    'List of Figures': 'vi',
    'List of Tables': 'vii',
    'List of Abbreviations': 'viii',

    # Chapters & Sections
    'Chapter 1': '13',
    '1.1': '13', '1.2': '13', '1.3': '13', '1.4': '14', '1.5': '14', '1.6': '14', '1.7': '15', '1.8': '15',

    'Chapter 2': '17',
    '2.1': '17', '2.2': '17', '2.3': '18', '2.4': '18', '2.5': '18', '2.6': '18', '2.7': '19',

    'Chapter 3': '20',
    '3.1': '20', '3.2': '20', '3.3': '21', '3.4': '21', '3.5': '22', '3.6': '22', '3.7': '23', '3.8': '23',

    'Chapter 4': '24',
    '4.1': '24', '4.2': '24', '4.3': '24', '4.4': '24', '4.5': '25', '4.6': '25', '4.7': '26',

    'Chapter 5': '28',
    '5.1': '28', '5.2': '28', '5.3': '29', '5.4': '30', '5.5': '30', '5.6': '30', '5.7': '31', '5.8': '31', '5.9': '31', '5.10': '31',

    'Chapter 6': '34',
    '6.1': '34', '6.2': '34', '6.3': '35', '6.4': '35', '6.5': '35', '6.6': '36', '6.7': '36',

    'Chapter 7': '37',
    '7.1': '37', '7.2': '37', '7.3': '38', '7.4': '38', '7.5': '39', '7.6': '39', '7.7': '39', '7.8': '40', '7.9': '40', '7.10': '40',

    'Chapter 8': '41',
    '8.1': '41', '8.2': '41', '8.3': '41', '8.4': '42', '8.5': '42', '8.6': '42', '8.7': '43', '8.8': '43', '8.9': '43', '8.10': '44',

    'Chapter 9': '45',
    '9.1': '45', '9.2': '45', '9.3': '45', '9.4': '47', '9.5': '47', '9.6': '47', '9.7': '48',

    'Chapter 10': '49',
    '10.1': '49', '10.2': '49', '10.3': '49', '10.4': '49', '10.5': '50', '10.6': '50',

    'Chapter 11': '51',
    '11.1': '51', '11.2': '51', '11.3': '51', '11.4': '52', '11.5': '52',

    'Chapter 12': '53',
    '12.1': '53', '12.2': '53', '12.3': '53', '12.4': '53', '12.5': '53', '12.6': '54', '12.7': '54',

    'Chapter 13': '55',
    '13.1': '55', '13.2': '55', '13.3': '55', '13.4': '55', '13.5': '56',

    'Chapter 14': '57',
    '14.1': '57', '14.2': '57', '14.3': '57', '14.4': '58',

    'Chapter 15': '59',
    '15.1': '59', '15.2': '59', '15.3': '59', '15.4': '59',

    # Appendices & End Chapters
    'Chapter A': '60',
    'Appendix A': '60',
    'Appendix B': '60',
    'Appendix C': '60',
    'Appendix D': '60',
    'Appendix E': '60',
    'Appendix F': '61',
    'Appendix G': '61',
    'Appendix H': '62',

    'Chapter B': '63',
    'Chapter C': '64'
}

# Update Table 3 (Table of Contents)
toc = doc.tables[3]
print(f"Updating Table 3 (TOC) with {len(toc.rows)} rows...")

for idx, row in enumerate(toc.rows):
    if idx == 0:
        continue
    col0 = row.cells[0].text.strip()
    col1 = row.cells[1].text.strip()
    
    page_val = ""
    if col0 in PAGE_MAP:
        page_val = PAGE_MAP[col0]
    elif col1 in PAGE_MAP:
        page_val = PAGE_MAP[col1]

    if page_val:
        is_bold = col0.startswith('Chapter') or col0.startswith('Appendix') or col1 == 'Preliminary Pages' or col1 == 'Appendices'
        set_cell(row.cells[2], page_val, bold=is_bold, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)

# Update List of Figures (Paragraph 45)
print("Updating List of Figures...")
set_para_text(doc.paragraphs[45],
"Figure 4.1: Overall Project Workflow & Agile Engineering Lifecycle ................................. 25\n"
"Figure 5.1: Final 3-Tier System Architecture & Cloud Infrastructure ................................. 28\n"
"Figure 5.2: End-to-End Maintenance & AR Repair Process Flow ......................................... 29\n"
"Figure 5.3: Major UML Class & Relational Database Schema Diagram ............................... 32\n"
"Figure 5.4: Major Interface & UI Wireframe Navigation Hierarchy ...................................... 32\n"
"Figure 7.1: Module 1 — Live AR Step Guidance with Pulsing Canvas Markers .................... 37\n"
"Figure 7.2: Module 2 — Multimodal AI Computer Vision Tool Verification Modal .................. 38\n"
"Figure 7.3: Module 3 — Offline-First Write-Ahead Sync Pipeline Architecture ....................... 39\n"
"Figure 9.1: Final Output 1 — Machine Specifications & Procedure Selection ........................... 45\n"
"Figure 9.2: Final Output 2 — Live AR Step Guidance with Hands-Free Voice Control .................. 46\n"
"Figure 9.3: Final Output 3 — Spare Parts Inventory & Low Stock Tracking ............................. 46\n"
"Figure 9.4: Final Output 4 — Supervisor Maintenance Analytics & MTTR Tracking ................... 47\n"
"Figure 10.1: Live Field Demonstration & End-to-End Operational Workflow ........................... 50\n"
"Figure A.1: Supporting Evidence — Plant QR Code Batch Generator & Print Matrix ................... 60\n"
"Figure A.2: Supporting Evidence — Fault Reporting & Photo Evidence Upload .......................... 61",
font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=4)

# Update List of Tables (Paragraph 49)
print("Updating List of Tables...")
set_para_text(doc.paragraphs[49],
"Table 1.1: Scope Matrix — Included, Excluded, and Operational Constraints ..................... 15\n"
"Table 2.1: Literature Review & Existing Approaches .................................................... 17\n"
"Table 2.2: Comparative Analysis of Maintenance Approaches ................................... 19\n"
"Table 3.1: Functional Requirements Specification .................................................... 20\n"
"Table 3.2: Non-Functional Requirements Specification ............................................. 21\n"
"Table 3.3: Hardware Requirements & Test Environments ........................................... 21\n"
"Table 3.4: Software Requirements & Framework Libraries ......................................... 22\n"
"Table 3.5: System Constraints & Operational Assumptions ........................................... 23\n"
"Table 4.1: Project Phases and Key Engineering Deliverables ......................................... 26\n"
"Table 4.2: Progress Report Milestone Consolidation .................................................... 26\n"
"Table 4.3: Project Execution Schedule & 16-Week Timeline ........................................... 27\n"
"Table 5.1: Module-Wise Architectural Design ............................................................. 30\n"
"Table 5.2: Relational Database Schema Design ............................................................ 33\n"
"Table 6.1: Programming Languages ............................................................................ 34\n"
"Table 6.2: Frameworks, Libraries, and APIs ................................................................ 34\n"
"Table 6.3: Database and Storage Technologies .............................................................. 35\n"
"Table 6.4: APIs and Third-Party Services ..................................................................... 35\n"
"Table 6.5: Development Tools and IDE ...................................................................... 36\n"
"Table 7.1: API & External Service Integration Specifications ...................................... 40\n"
"Table 8.1: Unit & Module Testing Test Cases ............................................................. 41\n"
"Table 8.2: Integration Testing Specifications ............................................................. 41\n"
"Table 8.3: Functional Testing Specifications ............................................................. 42\n"
"Table 8.4: System Testing & End-to-End Scenarios .................................................... 42\n"
"Table 8.5: Performance Testing Benchmarks ................................................................ 43\n"
"Table 8.6: User Acceptance Testing Across Physical Devices ........................................ 43\n"
"Table 8.7: Experimental Validation Benchmarks ........................................................ 44\n"
"Table 8.8: Defects Identification and Resolution .......................................................... 44\n"
"Table 9.1: Quantitative Performance Metrics & MTTR Reductions ............................... 47\n"
"Table 9.2: Comparative Benchmark with Legacy Maintenance Systems ........................ 48\n"
"Table 9.3: Objective-Wise Achievement Matrix ............................................................. 48\n"
"Table 10.1: Production Deployment Environment & Configuration ............................... 49\n"
"Table 11.1: Major Engineering Challenges & Solutions ............................................... 51\n"
"Table 11.2: Improvements Made During Development .................................................... 52\n"
"Table 12.1: Individual Contribution Matrix (Het Joshi) ................................................ 54\n"
"Table 13.1: Current System Limitations & Boundary Factors ........................................... 55\n"
"Table 14.1: Final Project Status Matrix ..................................................................... 58\n"
"Table F.1: Repository Structure & Core Modules .......................................................... 61\n"
"Table B.1: Final Hard-Bound Submission Checklist .................................................... 63",
font_size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=4)

print("Saving updated docx...")
doc.save(OUTPUT_PRO_DOCX)
print("DOCX updated successfully!")

# Export to PDF via clean Word COM
print("Exporting finalized PDF via Microsoft Word...")
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = False

try:
    doc_word = word.Documents.Open(os.path.abspath(OUTPUT_PRO_DOCX), ReadOnly=True)
    doc_word.SaveAs(os.path.abspath(OUTPUT_PRO_PDF), FileFormat=17) # 17 = wdFormatPDF
    doc_word.Close(SaveChanges=False)
    print(f"Final PDF exported successfully: {OUTPUT_PRO_PDF}")
finally:
    word.Quit()

# Copy to Downloads
try:
    shutil.copy2(OUTPUT_PRO_DOCX, OUTPUT_DL_DOCX)
    shutil.copy2(OUTPUT_PRO_PDF, OUTPUT_DL_PDF)
    print("Files successfully copied to Downloads folder!")
except Exception as e:
    print(f"Downloads copy notice: {e}")

print("ALL FINAL HARDBOUND DELIVERABLES COMPLETED PERFECTLY!")
