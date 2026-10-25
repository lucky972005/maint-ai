import docx

doc = docx.Document(r"C:\Users\hetjo\Downloads\PPSU_Final_Hard_Bound_Project_Report_Format.docx")
body = doc._body._element

ch1_idx = None
for i, el in enumerate(body):
    txt = "".join([t.text for t in el.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text])
    if txt.strip() == "CHAPTER 1":
        ch1_idx = i
        break

print(f"ch1_idx in template: {ch1_idx}")
for i in range(ch1_idx - 5, ch1_idx + 3):
    tag = body[i].tag.split('}')[-1]
    txt = "".join([t.text for t in body[i].findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text])
    print(f"  El {i}: <{tag}> '{txt[:50]}'")

print(f"\nTotal tables in template: {len(doc.tables)}")
t4 = doc.tables[4]
print(f"Table 4 has {len(t4.rows)} rows, {len(t4.columns)} cols")
for r in t4.rows:
    print("  Row:", [c.text.strip() for c in r.cells])
