import os
import pptx
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
import win32com.client

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(2))
tf = txBox.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = 'P I T C H   D E C K'
run.font.name = 'Century Gothic'
run.font.size = Pt(54)
run.font.bold = True

rPr = run._r.get_or_add_rPr()
outline_xml = parse_xml(r'<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="25400"><a:solidFill><a:srgbClr val="00F0FF"/></a:solidFill></a:ln>')
rPr.append(outline_xml)
nofill_xml = parse_xml(r'<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')
rPr.append(nofill_xml)

out_pptx = os.path.abspath('test_outline.pptx')
prs.save(out_pptx)
print('Saved pptx successfully!')

ppt_app = win32com.client.Dispatch('PowerPoint.Application')
pres = ppt_app.Presentations.Open(out_pptx, WithWindow=False)
pres.Slides[0].Export(os.path.abspath('test_outline.png'), 'PNG', 1920, 1080)
pres.Close()
print('Exported test_outline.png successfully!')
